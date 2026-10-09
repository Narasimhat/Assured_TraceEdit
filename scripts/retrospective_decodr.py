"""Run native TraceEdit on unambiguously linked, documented DECODR pairs.

This is tool agreement, not an independent accuracy benchmark. Unsupported HDR
experiments and ambiguous names stay in a review queue.
"""
import hashlib,io,json,re,sys,zipfile
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from batch import BatchRunner
from corpus import longpath

def run():
    app=Path(__file__).resolve().parents[1];runner=BatchRunner(app)
    byname={}
    for digest,r in runner.reads.items():
        for s in r['sources']:
            base=re.split(r'[\\/!]',s['path'])[-1].lower()
            byname.setdefault(base,set()).add(digest)
    catalog=json.loads((app/'public/catalog/summary.json').read_text(encoding='utf-8'))
    jobs=[];comparisons={};skips=[];seen=set()
    for group in catalog['projects']:
        entries=json.loads((app/'public/catalog'/f"{group['id']}.json").read_text(encoding='utf-8'))['files']
        for entry in entries:
            path=entry['path']
            if 'decodr' not in path.lower() or entry['name'].lower()!='results.xlsx':continue
            try:
                if entry['source']=='zip member':
                    with zipfile.ZipFile(longpath(runner.root/entry['archive_path'])) as z:raw=z.read(entry['member'])
                else:raw=longpath(runner.root/path).read_bytes()
                digest=hashlib.sha256(raw).hexdigest()
                if digest in seen:continue
                seen.add(digest)
                wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
                rows=list(wb.worksheets[0].iter_rows(max_row=200,max_col=10,values_only=True));wb.close()
                sample=str(rows[0][0]);meta={str(r[0]).strip():r[1:] for r in rows if r[0] is not None}
                control=str(meta['Wild-type (control) data:'][0]);guides=[str(x).upper() for x in meta['Guide sequences:'] if x]
                nuclease=meta.get('Nuclease:',('unknown',))[0]
                if nuclease not in ('Cas9','SpCas9'):raise ValueError('Unsupported or unspecified nuclease: '+str(nuclease))
                if len(guides) not in (1,2) or any(not re.fullmatch('[ACGT]{17,25}',g) for g in guides):raise ValueError('Unsupported guides')
                sh=byname.get(sample.lower(),set());ch=byname.get(control.lower(),set())
                if len(sh)!=1 or len(ch)!=1:raise ValueError(f'Names resolve to {len(sh)} sample and {len(ch)} control hashes')
                header=next(i for i,r in enumerate(rows) if r[0]=='Contribution')
                hdr=[r for r in rows[header+1:] if isinstance(r[0],(int,float)) and any(x not in (None,False,0,'No','False','None','N/A','-') for x in r[5:7])]
                if hdr:raise ValueError('HDR/partial HDR evidence requires donor model review')
                sid,cid=next(iter(sh)),next(iter(ch))
                if sid==cid:raise ValueError('Sample equals control bytes; needs separate negative-control review')
                job={'id':'decodr-'+digest[:16],'sample_sha256':sid,'control_sha256':cid,'settings':{'guides':guides},'pairing_evidence':{'kind':'saved_tool_export','path':path,'sample':sample,'control':control},'project':group['name']}
                jobs.append(job);comparisons[job['id']]={'sample':sample,'control':control,'source':path,'nuclease_recorded':nuclease,'decodr_r_squared':meta.get('R-squared:',(None,))[0],'decodr_ko_score_percent':meta.get('Knockout score:',(None,))[0]}
            except Exception as e:skips.append({'path':path,'reason':str(e)})
    (runner.out/'decodr-manifest.json').write_text(json.dumps({'jobs':jobs},indent=2),encoding='utf-8')
    print(json.dumps({'eligible_jobs':len(jobs),'review_records':len(skips)}),flush=True)
    results=runner.run(jobs,lambda n,total,r:print(json.dumps({'done':n,'total':total,'status':r['status']}),flush=True)) if jobs else {'results':[]}
    rows=[]
    for r in results['results']:
        row={**comparisons[r['id']],'status':r['status'],'seconds':r['seconds'],'error':r.get('error')}
        if 'result' in r:
            native=r['result'];row.update(native_indel_fraction=native['metrics']['indel_fraction'],native_r_squared=native['metrics']['r_squared'],warnings=native['warnings'])
            # Align the conventional KO-score proxy definition for comparison only.
            proxy=100*sum(o['fraction'] for o in native['outcomes'] if o['net_bp']%3 or abs(o['net_bp'])>=21)
            row['native_ko_proxy_percent']=proxy
            if isinstance(row['decodr_ko_score_percent'],(int,float)):
                row['absolute_ko_difference_pp']=abs(proxy-row['decodr_ko_score_percent'])
        rows.append(row)
    report={'eligible_pairs':len(jobs),'review_queue':skips,'results':rows,'interpretation':'Agreement with historical DECODR outputs, not ground truth. Explicit sample/control names were linked only when each resolved to one unique byte hash. Generic Cas9 exports are interpreted as SpCas9. KO scores are model-dependent proxies and do not establish functional knockout.'}
    (app/'public/corpus/retrospective.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps({'finished':len(rows),'errors':sum(r['status']=='error' for r in rows)}),flush=True)

if __name__=='__main__':run()
