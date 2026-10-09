"""Extract local evidence text and comparator tables; never promote them to truth labels."""
import csv, hashlib, io, json, re, sqlite3, sys, time, zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET
from html.parser import HTMLParser
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from corpus import longpath
from openpyxl import load_workbook
from pypdf import PdfReader

APP=Path(__file__).resolve().parents[1]
OUT=APP/'data/evidence';OUT.mkdir(parents=True,exist_ok=True)
TEXT=OUT/'text';TEXT.mkdir(exist_ok=True)
DB=sqlite3.connect(OUT/'evidence.sqlite')
DB.execute('CREATE TABLE IF NOT EXISTS documents(hash TEXT PRIMARY KEY, data TEXT)')
EXTS={'.pdf','.docx','.xlsx','.csv','.tsv','.txt','.md','.html','.htm','.json','.gb','.gbk','.fasta','.fa','.seq'}
MAX_BYTES=30_000_000

class TextHTML(HTMLParser):
    def __init__(self):super().__init__();self.parts=[];self.skip=0
    def handle_starttag(self,tag,attrs):
        if tag in ('script','style'):self.skip+=1
    def handle_endtag(self,tag):
        if tag in ('script','style'):self.skip=max(0,self.skip-1)
    def handle_data(self,data):
        if not self.skip:self.parts.append(data)

def extract(raw, ext):
    tables=[];limits=[]
    if ext=='.xlsx':
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        chunks=[]
        for sheet in wb:
            if (sheet.max_row or 0)>10000 or (sheet.max_column or 0)>300: limits.append('Sheet truncated to 10000 rows / 300 columns')
            rows=[]
            for i,row in enumerate(sheet.iter_rows(max_row=min(sheet.max_row or 10000,10000),max_col=min(sheet.max_column or 300,300),values_only=True),1):
                row=list(row)
                while row and row[-1] is None:row.pop()
                if row:rows.append({'row':i,'cells':row})
            tables.append({'sheet':sheet.title,'rows':rows})
            chunks.append(sheet.title+'\n'+'\n'.join(str(r['row'])+'\t'+'\t'.join('' if c is None else str(c) for c in r['cells']) for r in rows))
        wb.close();text='\n'.join(chunks)
    elif ext=='.docx':
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            entry=z.getinfo('word/document.xml')
            if entry.file_size>30_000_000:raise ValueError('DOCX XML exceeds size limit')
            tree=ET.fromstring(z.read(entry))
            text='\n'.join(''.join(p.itertext()) for p in tree.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'))
    elif ext=='.pdf':
        reader=PdfReader(io.BytesIO(raw))
        if len(reader.pages)>100:limits.append('PDF truncated at 100 pages')
        text='\n'.join(f'[Page {i+1}]\n'+(page.extract_text() or '') for i,page in enumerate(reader.pages[:100]))
        if len(text.strip())<50:limits.append('Little extracted text; visual/OCR review needed')
    else:
        text=raw.decode('utf-8-sig',errors='replace')
        if ext in ('.html','.htm'):
            p=TextHTML();p.feed(text);text='\n'.join(p.parts)
        elif ext in ('.csv','.tsv'):
            rows=list(csv.reader(io.StringIO(text),delimiter='\t' if ext=='.tsv' else ','))
            tables=[{'sheet':'text','rows':[{'row':i+1,'cells':row} for i,row in enumerate(rows) if any(row)]}]
    if len(text)>1_000_000:limits.append('Text truncated at 1 million characters');text=text[:1_000_000]
    return {'text':text,'tables':tables,'limits':limits}

def run():
    catalog=json.loads((APP/'public/catalog/summary.json').read_text(encoding='utf-8'))
    root=Path(catalog['source_root']);records=[];comparators=[];errors=[];temporary_files=[];counts=Counter();n=0
    for group in catalog['projects']:
        archives={}
        entries=json.loads((APP/'public/catalog'/f"{group['id']}.json").read_text(encoding='utf-8'))['files']
        for entry in entries:
            ext=Path(entry['name']).suffix.lower()
            if ext not in EXTS:continue
            path=entry['path'];n+=1
            if entry['name'].startswith('~$'):
                temporary_files.append({'path':path,'reason':'Office lock file, not a source document'});continue
            try:
                if entry['bytes']>MAX_BYTES:raise ValueError('Document exceeds 30 MB extraction limit')
                if entry['source']=='zip member':
                    archive=entry['archive_path']
                    if archive not in archives:archives[archive]=zipfile.ZipFile(longpath(root/archive))
                    with archives[archive].open(entry['member']) as stream:raw=stream.read(MAX_BYTES+1)
                else:
                    with longpath(root/path).open('rb') as stream:raw=stream.read(MAX_BYTES+1)
                if len(raw)>MAX_BYTES:raise ValueError('Document exceeds size limit')
                digest=hashlib.sha256(raw).hexdigest()
                cached=DB.execute('SELECT data FROM documents WHERE hash=?',(digest,)).fetchone()
                if cached:data=json.loads(cached[0])
                else:
                    data=extract(raw,ext)
                    DB.execute('INSERT INTO documents VALUES(?,?)',(digest,json.dumps(data,default=str)))
                    (TEXT/(digest+'.txt')).write_text(data['text'],encoding='utf-8')
                text=data['text'];lower=path.lower();kind='document'
                # Filenames/context classify evidence type, never outcome correctness.
                if 'project plan' in lower or 'design_report' in lower:kind='planned_design'
                elif re.search(r'(^|[\\/!_ -])ice([_ .\\/!-]|$)|decod[er]+',lower):kind='tool_comparison'
                elif re.search(r'report|verification|analysis|summary|final',lower):kind='review_candidate'
                elif ext in ('.fa','.fasta','.gb','.gbk','.seq'):kind='sequence_text'
                counts[kind]+=1
                records.append({'path':path,'project':group['name'],'year':group['year'],'sha256':digest,'kind':kind,'characters':len(text),'limits':data['limits'],'excerpt':text[:900], 'truth_status':'unreviewed'})
                if kind=='tool_comparison':
                    for table in data['tables']:
                        rows=table['rows']
                        for j,row in enumerate(rows):
                            cells=[str(x or '').strip() for x in row['cells']]
                            if 'Label' in cells and 'ICE' in cells:
                                for r in rows[j+1:]:
                                    values=dict(zip(cells,r['cells']))
                                    if values.get('Label'):comparators.append({'tool':'ICE','path':path,'sheet':table['sheet'],'row':r['row'],'sample_label':str(values['Label']),'raw_metrics':values,'truth_status':'comparator_only'})
                                break
                            if cells and cells[0]=='Contribution' and 'Net Indel' in cells:
                                sample=str(rows[0]['cells'][0])
                                alleles=[{'row':r['row'],'values':dict(zip(cells[:9],r['cells'][:9]))} for r in rows[j+1:] if r['cells'] and isinstance(r['cells'][0],(int,float))]
                                comparators.append({'tool':'DECODR','path':path,'sheet':table['sheet'],'row':row['row'],'sample_label':sample,'alleles':alleles,'truth_status':'comparator_only'})
                                break
            except Exception as e:errors.append({'path':path,'error':str(e)[:300]})
            if n%100==0:DB.commit();print(json.dumps({'processed':n,'extracted':len(records),'comparators':len(comparators)}),flush=True)
        for archive in archives.values():archive.close()
    DB.commit()
    # Preserve ambiguous mappings: identical copies collapse; different hashes never do.
    reads=json.loads((APP/'public/corpus/reads.json').read_text(encoding='utf-8'))
    byname={}
    for r in reads:
        for s in r['sources']:
            key=re.split(r'[\\/!]',s['path'])[-1].lower();key=re.sub(r'\.(ab1|abi)$','',key)
            byname.setdefault(key,set()).add(r['sha256'])
    for c in comparators:
        key=re.sub(r'\.(ab1|abi)$','',c['sample_label'].strip().lower())
        c['candidate_read_hashes']=sorted(byname.get(key,set()))
        c['link_status']='exact_name_unique_hash' if len(c['candidate_read_hashes'])==1 else 'ambiguous' if c['candidate_read_hashes'] else 'unmatched'
    out=APP/'public/corpus'
    summary={'created_utc':time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime()),'attempted':n,'extracted_occurrences':len(records),'unique_documents':len(set(x['sha256'] for x in records)),'kinds':dict(counts),'comparator_rows':len(comparators),'comparator_tools':dict(Counter(c['tool'] for c in comparators)),'link_status':dict(Counter(c['link_status'] for c in comparators)),'errors':errors,'coverage':'Text/tables extracted from supported physical documents and first-level ZIP members. Images, legacy binary documents, nested archive documents and PDF figures need separate review. Formula values are saved workbook caches, not recalculated. No manual labels verified.'}
    summary['temporary_files']=temporary_files
    for name,data in [('evidence.json',records),('comparators.json',comparators),('evidence-summary.json',summary)]:
        tmp=out/(name+'.tmp');tmp.write_text(json.dumps(data,ensure_ascii=False,default=str),encoding='utf-8');tmp.replace(out/name)
    print(json.dumps(summary),flush=True)

if __name__=='__main__':run()
